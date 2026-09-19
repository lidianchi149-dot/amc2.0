from __future__ import annotations

import hashlib
import json
import re
from zipfile import BadZipFile
from dataclasses import asdict, dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from sqlalchemy import text

from .database import engine


EXPECTED_HEADERS = (
    "中文名称", "英文名称", "化学式", "分子式", "CAS号", "沸点(℃)", "摩尔质量(g/mol)",
)


@dataclass
class PollutantRow:
    source_row: int
    name_cn: str
    name_en: str | None
    formula: str | None
    molecular_formula: str | None
    cas_no: str
    boiling_point: Decimal | None
    molar_mass: Decimal
    aliases: list[str]


def normalize_cas(value: Any) -> str:
    """Normalize common Excel punctuation/zero-padding issues and verify CAS checksum."""
    groups = re.findall(r"\d+", str(value or ""))
    if len(groups) != 3:
        raise ValueError("CAS 号必须包含三段数字")
    first = groups[0].lstrip("0") or "0"
    middle = groups[1].zfill(2)
    check = groups[2]
    if len(check) > 1 and set(check[:-1]) <= {"0"}:
        check = check[-1]
    if not 2 <= len(first) <= 7 or len(middle) != 2 or len(check) != 1:
        raise ValueError("CAS 号格式不正确")
    digits = first + middle
    checksum = sum(int(digit) * weight for weight, digit in enumerate(reversed(digits), 1)) % 10
    if checksum != int(check):
        raise ValueError("CAS 号校验位不正确")
    return f"{first}-{middle}-{check}"


def read_pollutant_workbook(file_path: str | Path) -> tuple[list[PollutantRow], dict[str, Any]]:
    path = Path(file_path).expanduser().resolve()
    if not path.is_file():
        raise FileNotFoundError(f"污染物 Excel 不存在：{path}")
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, OSError) as exc:
        raise ValueError("文件不是可读取的 XLSX 工作簿") from exc
    try:
        sheet = workbook.active
        headers = tuple(_text(sheet.cell(1, column).value) for column in range(2, 9))
        if headers != EXPECTED_HEADERS:
            raise ValueError(f"污染物 Excel 表头不匹配，实际表头：{headers}")

        by_cas: dict[str, PollutantRow] = {}
        warnings: list[dict[str, Any]] = []
        total_rows = 0
        duplicate_rows = 0
        for source_row, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            sequence = values[0]
            if not isinstance(sequence, (int, float)):
                continue
            total_rows += 1
            name_cn = _required_text(values[1], source_row, "中文名称")
            raw_cas = _required_text(values[5], source_row, "CAS号")
            cas_no = normalize_cas(raw_cas)
            molar_mass = _positive_decimal(values[7], source_row, "摩尔质量")
            boiling_point = _decimal_or_none(values[6], source_row, "沸点")
            item = PollutantRow(
                source_row=source_row,
                name_cn=name_cn,
                name_en=_text(values[2]) or None,
                formula=_text(values[3]) or None,
                molecular_formula=_text(values[4]) or None,
                cas_no=cas_no,
                boiling_point=boiling_point,
                molar_mass=molar_mass,
                aliases=[],
            )
            if _text(raw_cas) != cas_no:
                warnings.append({
                    "row": source_row, "field": "CAS号", "type": "normalized",
                    "source": _text(raw_cas), "value": cas_no,
                })
            existing = by_cas.get(cas_no)
            if existing is None:
                by_cas[cas_no] = item
                continue
            duplicate_rows += 1
            for alias in (item.name_cn, item.name_en):
                if alias and alias not in {existing.name_cn, existing.name_en} and alias not in existing.aliases:
                    existing.aliases.append(alias)
            warnings.append({
                "row": source_row, "field": "CAS号", "type": "duplicate_merged",
                "value": cas_no, "primaryRow": existing.source_row,
            })

        metadata = {
            "file": path,
            "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
            "totalRows": total_rows,
            "uniqueRows": len(by_cas),
            "duplicateRows": duplicate_rows,
            "warnings": warnings,
        }
        return list(by_cas.values()), metadata
    finally:
        workbook.close()


def import_pollutants(file_path: str | Path, username: str, *, file_name: str | None = None) -> dict[str, Any]:
    rows, metadata = read_pollutant_workbook(file_path)
    path: Path = metadata["file"]
    display_name = Path(file_name).name if file_name else path.name
    inserted = 0
    updated = 0
    with engine.begin() as connection:
        user = connection.execute(text("""
            SELECT id, username FROM sys_user WHERE username=:username AND status='active'
        """), {"username": username}).mappings().first()
        if user is None:
            raise ValueError(f"有效用户不存在：{username}")
        for row in rows:
            existing_id = connection.execute(
                text("SELECT id FROM pollutant WHERE cas_no=:cas_no"), {"cas_no": row.cas_no}
            ).scalar_one_or_none()
            values = {
                **asdict(row),
                "aliases": json.dumps(row.aliases, ensure_ascii=False),
                "created_by": user["id"],
            }
            values.pop("source_row")
            if existing_id is None:
                connection.execute(text("""
                    INSERT INTO pollutant (
                      name_cn, name_en, aliases, formula, molecular_formula, cas_no,
                      boiling_point, molar_mass, status, created_by
                    ) VALUES (
                      :name_cn, :name_en, :aliases, :formula, :molecular_formula, :cas_no,
                      :boiling_point, :molar_mass, 'active', :created_by
                    )
                """), values)
                inserted += 1
            else:
                values["id"] = existing_id
                connection.execute(text("""
                    UPDATE pollutant SET name_cn=:name_cn, name_en=:name_en, aliases=:aliases,
                      formula=:formula, molecular_formula=:molecular_formula, cas_no=:cas_no,
                      boiling_point=:boiling_point, molar_mass=:molar_mass, status='active'
                    WHERE id=:id
                """), values)
                updated += 1
        detail = {**metadata, "file": display_name, "inserted": inserted, "updated": updated}
        result = connection.execute(text("""
            INSERT INTO import_record (
              import_type, template_version, file_name, file_hash, stage, total_rows,
              valid_rows, warning_rows, error_rows, duplicate_rows, detail_json,
              imported_by, finished_at
            ) VALUES (
              'pollutant', 'AMC明细汇总-V1', :file_name, :file_hash, 'imported', :total_rows,
              :valid_rows, :warning_rows, 0, :duplicate_rows, :detail_json, :imported_by, NOW()
            )
        """), {
            "file_name": display_name,
            "file_hash": metadata["sha256"],
            "total_rows": metadata["totalRows"],
            "valid_rows": metadata["uniqueRows"],
            "warning_rows": len(metadata["warnings"]),
            "duplicate_rows": metadata["duplicateRows"],
            "detail_json": json.dumps(detail, ensure_ascii=False, default=str),
            "imported_by": user["id"],
        })
        import_id = int(result.lastrowid)
        connection.execute(text("""
            INSERT INTO operation_log (
              user_id, username_snapshot, action, target_type, target_id, target,
              detail, result_status
            ) VALUES (
              :user_id, :username, 'import_pollutants', 'import_record', :target_id,
              :target, :detail, 'success'
            )
        """), {
            "user_id": user["id"], "username": user["username"], "target_id": import_id,
            "target": display_name,
            "detail": f"导入 {metadata['uniqueRows']} 个唯一污染物，合并 {metadata['duplicateRows']} 个重复项",
        })
    return {
        "importId": import_id,
        "file": display_name,
        "totalRows": metadata["totalRows"],
        "uniqueRows": metadata["uniqueRows"],
        "duplicateRows": metadata["duplicateRows"],
        "warningRows": len(metadata["warnings"]),
        "inserted": inserted,
        "updated": updated,
    }


def _text(value: Any) -> str:
    return str(value or "").replace("\u00a0", " ").strip()


def _required_text(value: Any, row: int, field: str) -> str:
    result = _text(value)
    if not result:
        raise ValueError(f"第 {row} 行缺少{field}")
    return result


def _positive_decimal(value: Any, row: int, field: str) -> Decimal:
    result = _decimal_or_none(value, row, field)
    if result is None or result <= 0:
        raise ValueError(f"第 {row} 行{field}必须大于 0")
    return result


def _decimal_or_none(value: Any, row: int, field: str) -> Decimal | None:
    if value in (None, ""):
        return None
    try:
        return Decimal(str(value))
    except (InvalidOperation, ValueError) as exc:
        raise ValueError(f"第 {row} 行{field}不是有效数字") from exc
